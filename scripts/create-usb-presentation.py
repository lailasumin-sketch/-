"""Create five editable presentation slides from actual USB screenshots."""
import argparse,json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR,MSO_AUTO_SIZE
from PIL import Image

ap=argparse.ArgumentParser();ap.add_argument('folder',type=Path);args=ap.parse_args();root=args.folder.resolve();out=root/'발표자료'
data=json.loads((root/'data/bundled-data.json').read_text());count_tour=len(data['catalogue']['places']);count_food=sum(p['category']=='맛집' for p in data['dining']);count_cafe=sum(p['category']=='카페' for p in data['dining'])
P=Presentation();P.slide_width=Inches(13.333);P.slide_height=Inches(7.5)
C={'bg':'F7F9FA','white':'FFFFFF','navy':'19374A','green':'356B51','light':'EAF1EC','orange':'D98851','muted':'637583','line':'DCE4E8'}
FONT='맑은 고딕'
def box(s,x,y,w,h,fill,line=None,round=False):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(C.get(fill,fill));sh.line.fill.background() if not line else None
 if line:sh.line.color.rgb=RGBColor.from_string(C.get(line,line))
 for effect in sh._element.xpath('./p:style/a:effectRef'):effect.set('idx','0')
 if round:
  try:sh.adjustments[0]=0.12
  except Exception:pass
 return sh

def text(s,x,y,w,h,value,size=18,color='navy',bold=False):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE
 tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(value.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(C.get(color,color));p.space_after=Pt(7)
 return sh

def base(n,label,title,sub):
 s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(C['bg'])
 box(s,.55,.35,.38,.38,'green',round=True);text(s,.62,.39,.25,.25,'의',12,'white',True)
 text(s,1.05,.38,5,.3,'의동이 PICK  |  바이브코딩 대회',12,'green',True)
 text(s,10.3,.38,2.5,.3,label,12,'muted')
 text(s,.62,1.02,12.05,.68,title,29,'navy',True);text(s,.65,1.85,12,.53,sub,15,'muted')
 box(s,.62,7.02,12.07,.015,'line');text(s,.65,7.13,10.8,.2,'CSV 기반 오프라인 체험판 · 가상/미검증 자료 포함 · 의성군 공식 서비스 아님',9,'muted');text(s,11.8,7.09,.9,.28,f'{n:02} / 05',11,'green',True)
 return s

def image(s,path,x,y,w,h):
 iw,ih=Image.open(path).size;scale=min(w/iw,h/ih);ww,hh=iw*scale,ih*scale
 return s.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))

notes=[]
s=base(1,'01  문제 정의','흩어진 여행 정보를, 나에게 맞는 일정으로','해결하고자 하는 문제: 정보 탐색 → 조건 비교 → 일정 구성의 번거로움')
for x,num,title,body in [(.65,'01','정보를 찾는 부담','관광지·맛집·체험 정보를\n각각 찾아 비교해야 하는 상황'),(4.83,'02','조건마다 다른 선택','동행·시간·관심사에 따라\n필요한 여행 코스가 달라짐'),(9.01,'03','정보의 근거 확인','가상·미확인 자료와\n계산한 조건을 구분할 필요')]:
 box(s,x,2.8,3.68,2.62,'white','line',True);text(s,x+.25,3.05,.7,.48,num,23,'orange',True);text(s,x+.25,3.75,3.2,.46,title,21,'navy',True);text(s,x+.25,4.42,3.2,.75,body,15,'muted')
box(s,.65,5.83,12.04,.66,'light',round=True);text(s,.92,6.0,11.5,.3,'제안: 선택한 조건과 데이터 근거가 보이는 맞춤형 의성 여행 로드맵',18,'green',True)
notes.append('안녕하세요. 의동이 PICK은 의성 여행을 준비할 때 정보를 찾고 비교한 뒤 일정을 직접 조합해야 하는 번거로움을 줄이기 위한 체험용 서비스입니다. 같은 의성을 방문하더라도 가족 여행과 연인 여행, 반나절과 하루 여행의 조건은 서로 다릅니다. 그래서 장소를 많이 나열하는 것보다 사용자가 선택한 조건을 일정으로 연결하는 데 집중했습니다. 또 가상 데이터나 확인되지 않은 운영정보가 실제 정보처럼 보이지 않도록 구분했습니다. 이 문제 정의는 서비스가 해결하려는 상황이며, 별도의 이용자 설문이나 효과 측정 결과를 제시하는 것은 아닙니다.')

s=base(2,'02  서비스 소개','의동이 PICK — 취향은 내가, 의성 여행은 의동이가','조건 선택 또는 문장 입력으로 후보를 비교하고 시간순 일정표를 생성합니다.')
box(s,.65,2.6,7.6,3.97,'white','line',True);image(s,out/'화면_메인.png',.75,2.68,7.4,3.8)
for y,num,title,body in [(2.72,'1','여행 조건 입력','기간·동행·관심사·권역\n시작시간·활동시간·여행 속도'),(4.0,'2','데이터와 조건 비교','체류시간·가상 업소 영업시간\n휴무·예산·주차 조건 반영'),(5.28,'3','시간표와 추천 이유','방문·이동·대기를 순서대로\n초록 체크로 확인된 이유 표시')]:
 box(s,8.65,y,.38,.38,'green',round=True);text(s,8.76,y+.04,.2,.22,num,12,'white',True);text(s,9.2,y-.03,3.3,.4,title,19,'navy',True);text(s,9.2,y+.44,3.1,.64,body,13,'muted')
notes.append('서비스 이름은 의동이 PICK이며, 취향은 내가 고르고 의성 여행의 일정 초안은 의동이가 구성한다는 의미입니다. 왼쪽은 실제 USB 체험판 화면입니다. 사용자는 여행 기간, 동행자, 관심사와 권역을 고르고 시작시간과 활동 가능시간을 입력합니다. 자연어 입력도 기간과 관심사 등 지원하는 키워드를 읽어 같은 선택 조건에 반영합니다. 추천은 장소 데이터와 시간 제약을 비교하는 규칙으로 동작합니다. 결과에는 시간순 방문 일정과 간결한 추천 이유를 보여줍니다. 현재 구현은 외부 AI나 실시간 지도 API 호출에 의존하지 않으며, 데이터가 없는 조건은 확인이 필요하다고 표시합니다.')

s=base(3,'03  CSV 데이터 활용','파일로 관리하고, 검색 결과로 시각화합니다','CSV 원본 보존 · 한글 UTF-8 BOM 확인 · 실행 시 외부 데이터베이스 연결 없음')
for x,value,label in [(.7,count_tour,'관광 조사 후보 / 별도 JSON'),(3.17,count_food,'가상 맛집 / CSV'),(5.64,count_cafe,'가상 카페 / CSV')]:
 box(s,x,2.65,2.22,1.1,'white','line',True);text(s,x+.18,2.78,1.9,.46,str(value)+'곳',25,'green',True);text(s,x+.18,3.36,1.95,.23,label,10,'muted')
box(s,8.35,2.65,4.25,1.1,'light',round=True);text(s,8.58,2.83,3.85,.8,'가상 업소 CSV 30행 · 24필드\n메뉴·예산·체류·영업·휴무 등',15,'green',True)
text(s,.75,4.03,11.9,.43,'가상 업소 CSV → JSON 변환·검사 → JavaScript 내장 → 검색·필터·차트·일정',18,'navy',True)
def compact_chart(x,title,items):
 box(s,x,4.6,3.82,1.72,'white','line',True);text(s,x+.16,4.73,3.5,.25,title,12,'green',True)
 maximum=max(1,max(v for _,v in items))
 for j,(label,value) in enumerate(items):
  y=5.13+j*.255;text(s,x+.16,y,1.65,.2,label,10,'muted');box(s,x+1.88,y+.055,1.27,.07,'light');
  if value:box(s,x+1.88,y+.055,1.27*value/maximum,.07,'green')
  text(s,x+3.28,y-.015,.42,.23,str(value),11,'navy',True)
compact_chart(.75,'자료 분류', [('관광 조사 후보',count_tour),('가상 맛집',count_food),('가상 카페',count_cafe)])
regions={k:0 for k in data['catalogue']['meta']['regions']}
for place in data['catalogue']['places']+data['dining']:regions[place['region']]+=1
compact_chart(4.8,'권역별 레코드',[(data['catalogue']['meta']['regions'][k].removesuffix('권'),v) for k,v in regions.items()])
compact_chart(8.85,'가상 업소 1인 예산',[('1만원 이하',sum(p['price']<=10000 for p in data['dining'])),('1만~2만원',sum(10000<p['price']<=20000 for p in data['dining'])),('2만원 초과',sum(p['price']>20000 for p in data['dining']))])
text(s,.75,6.55,12,.2,'사용자가 승인한 기존 가상 CSV 사용. 주최 측 제공 파일로 확인된 자료는 없으며, 표시는 제출 데이터의 레코드 수이며 성과 통계가 아닙니다.',10,'muted')
notes.append('데이터는 관광 조사 후보 35곳과 가상 맛집 18곳, 가상 카페 12곳으로 구성했습니다. 맛집과 카페는 기존 엑셀과 같은 CSV 30행, 24개 필드를 사용하며 메뉴, 일인 예산, 동행 유형, 체류시간, 영업시간과 휴무 등의 조건이 있습니다. 관광 후보는 별도 JSON을 재사용하고 조사용 CSV도 함께 제출합니다. 두 CSV 모두 UTF-8 BOM 한글 인코딩을 확인했습니다. USB 빌드 시 데이터를 JavaScript에 내장해 파일 직접 실행에서 CSV 접근 요청이 필요 없도록 만들었습니다. 분류, 권역, 예산 구간을 차트로 요약합니다. 슬라이드의 차트는 제출 데이터 전체 분포이며, 웹 화면의 차트는 검색 결과에 따라 갱신됩니다. 수치는 방문 실적이나 매출이 아닌 레코드 수이고, 주최 측이 제공한 실제 데이터라고 설명하지 않습니다.')

s=base(4,'04  실제 화면','날짜별 여행 카드와 추천 이유를 확인합니다','실제 실행한 USB 체험판 화면 · 가상 데이터 표시와 미확인 정보 안내 유지')
box(s,.65,2.55,8.35,4.14,'white','line',True)
image(s,out/'화면_카드.png',.75,2.62,8.15,4.0)
for y,title,body in [(2.8,'시간순 일정 카드','방문 순서·시각·유형별 카드\n날짜 선택과 조건 재수정'),(4.07,'추천 이유 표시','동행·관심사·예산·권역 등\n실제 일치·계산한 이유만 설명'),(5.32,'자료 구분 유지','합산 가능한 1인 금액과 제외 항목\n미확인 위치는 지도 연결 제외')]:
 text(s,9.35,y,3.3,.38,'✓ '+title,18,'green',True);text(s,9.35,y+.51,3.25,.7,body,13,'muted')
notes.append('이 화면은 발표용으로 그린 목업이 아니라 실제 USB 패키지를 실행해 캡처한 결과입니다. 예시에서는 가족 여행과 여섯 시간의 활동 예산을 선택했습니다. 등록된 체류시간과 사용자가 입력한 이동 여유를 합산해 시간 예산에 맞춥니다. 이 이동 여유는 실제 도로 이동시간이 아닙니다. 유형 색상과 방문 번호로 순서를 읽고 날짜별 버튼으로 확인할 수 있습니다. 각 장소의 초록 체크에는 동행, 관심사, 예산과 권역 등 데이터와 일치한 조건을 보여줍니다. 등록된 일인 금액만 합산하고 교통비와 미등록 항목은 제외합니다. 점심시간 이유는 실제 배치한 시간대가 맞을 때만 붙습니다. 이동시간이 없는 구간은 숫자를 만들지 않고 확인 필요로 표시하며, 그 이후 시각은 확정이 아닌 최소 시각의 임시 일정으로 안내합니다. 검색과 필터를 바꾼 뒤 같은 방식으로 결과를 비교할 수 있습니다.')

s=base(5,'05  군민 혜택','군민의 여행 준비와 지역 자원의 연결을 돕습니다','다음 내용은 기대효과이며 실제 방문·매출·만족도 개선은 아직 측정하지 않았습니다.')
for x,title,body in [(.65,'군민·방문객','조건별 후보와 일정 초안을\n한 화면에서 비교하는 데 도움'),(4.83,'지역 상권','맛집·카페·특산물 판매 후보를\n추천 일정에 연결하는 구조'),(9.01,'운영·행정','CSV 기반 자료 관리와\n추천 근거·미확인 정보 구분')]:
 box(s,x,2.75,3.68,2.42,'white','line',True);text(s,x+.25,3.1,3.18,.45,title,23,'green',True);text(s,x+.25,3.94,3.18,.9,body,16,'muted')
box(s,.65,5.65,12.04,.95,'light',round=True);text(s,.9,5.82,11.55,.58,'현재: CSV 기반 오프라인 체험판\n다음: 실제 업소·운영 검증, 이동자료 확보, 군민 대상 사용성 평가',15,'green',True)
notes.append('군민과 방문객에게 기대하는 첫 번째 혜택은 조건별 후보와 일정 초안을 한 화면에서 비교할 수 있다는 점입니다. 지역 상권에는 맛집, 카페와 특산물 판매 후보를 관광 일정 안에 연결할 수 있는 구조를 제시합니다. 화면의 상권 연계 수는 추천된 데이터의 수이며 실제 매출이나 방문 성과가 아닙니다. 운영 측면에서는 CSV 자료를 관리하고 추천 근거와 미확인 항목을 구분할 수 있습니다. 현재 제출 버전은 가상 데이터가 포함된 오프라인 체험판입니다. 실제 업소와 운영 정보를 검증하고 이동자료를 확보한 뒤 군민 대상 사용성 평가로 효과를 확인하는 것이 다음 단계입니다. 감사합니다.')
for slide,note in zip(P.slides,notes):slide.notes_slide.notes_text_frame.text=note
P.core_properties.title='의동이 PICK | 5분 발표';P.core_properties.subject='실제 구현 기능과 가상 데이터 범위를 설명하는 USB 체험판 발표';P.core_properties.author='의동이 PICK';P.core_properties.keywords='의성, CSV, USB, 오프라인, 여행, 체험판'
file=out/'의동이_PICK_5분발표.pptx';P.save(file)
script='의동이 PICK 5분 발표 대본\n페이지당 약 1분을 목표로 한 발표자 노트입니다. 실제 발화 시간은 발표자가 연습해 조정하세요.\n\n'
for i,note in enumerate(notes,1):script+=f'{i}. {[("문제 정의"),("서비스 소개"),("CSV 데이터 활용"),("실제 화면"),("군민 혜택")][i-1]} ({i-1}:00–{i}:00)\n{note}\n\n'
(out/'5분_발표대본.txt').write_text(script,encoding='utf-8-sig')
print('Created',file,'slides',len(P.slides),'speaker characters',sum(map(len,notes)))
